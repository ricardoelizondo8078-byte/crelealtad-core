import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from src.normalization import normalize, normalize_postal_code, normalize_number
from src.utils import key_for, legacy_key_for
from src.consolidation import consolidate, needs_second_pass, score_result, OFFICIAL_COLUMNS, GRANULARITY_POINTS, _granularity_key
from src import processor
from src.cache import get, put
from src.sepomex import best_match, load_catalog
from src.workbook import GOOGLE_HEADERS, ensure_headers, write_row

def validation_payload(complete=True):
    return {"result": {"verdict": {"addressComplete": complete, "validationGranularity": "PREMISE",
                                     "geocodeGranularity": "PREMISE", "hasUnconfirmedComponents": False},
                       "address": {"formattedAddress": "Av. Juárez 100, Centro, Monterrey, N.L. 64000, México",
                                   "postalAddress": {"locality": "Monterrey", "postalCode": "64000", "administrativeArea": "Nuevo León"},
                                   "addressComponents": [
                                       {"componentType": "route", "componentName": {"text": "Avenida Juárez"}, "confirmationLevel": "CONFIRMED"},
                                       {"componentType": "street_number", "componentName": {"text": "100"}, "confirmationLevel": "CONFIRMED"},
                                       {"componentType": "neighborhood", "componentName": {"text": "Monterrey Centro"}, "confirmationLevel": "CONFIRMED"},
                                       {"componentType": "locality", "componentName": {"text": "Monterrey"}, "confirmationLevel": "CONFIRMED"},
                                       {"componentType": "postal_code", "componentName": {"text": "64000"}, "confirmationLevel": "CONFIRMED"},
                                       {"componentType": "administrative_area_level_1", "componentName": {"text": "Nuevo León"}, "confirmationLevel": "CONFIRMED"}]},
                       "geocode": {"location": {"latitude": 25.6866, "longitude": -100.3161}, "placeId": "place-1"}}}

class V5Tests(unittest.TestCase):
    def test_granularity_points_preserve_underscores(self):
        expected = {"SUB_PREMISE": 18, "PREMISE": 17, "PREMISE_PROXIMITY": 14,
                    "BLOCK": 10, "ROUTE": 7, "OTHER": 2}
        for value, points in expected.items():
            self.assertEqual(GRANULARITY_POINTS[_granularity_key(value)], points)
        self.assertEqual(_granularity_key("sub premise"), "SUB_PREMISE")
        self.assertEqual(_granularity_key("premise-proximity"), "PREMISE_PROXIMITY")

    def test_sepomex_loader_skips_legal_preamble(self):
        content = (
            "El Catálogo Nacional de Códigos Postales es elaborado por Correos de México.\n"
            "d_codigo|d_asenta|d_tipo_asenta|D_mnpio|d_estado\n"
            "64000|Monterrey Centro|Colonia|Monterrey|Nuevo León\n"
        )
        with tempfile.TemporaryDirectory() as temp:
            catalog_path = Path(temp) / "SEPOMEX.txt"
            catalog_path.write_text(content + (" " * 10_000), encoding="utf-8")
            with patch("src.sepomex.CATALOG_DIR", Path(temp)):
                rows = load_catalog()
        self.assertGreater(len(rows), 0)
        self.assertEqual(rows[0]["d_codigo"], "64000")
        self.assertEqual(rows[0]["d_asenta"], "Monterrey Centro")
        self.assertEqual(rows[0]["D_mnpio"], "Monterrey")
        self.assertEqual(rows[0]["d_estado"], "Nuevo León")

    def test_normalization_rules(self):
        self.assertEqual(normalize("Av. Juárez # 100"), "AVENIDA JUAREZ NUMERO 100")
        self.assertEqual(normalize_postal_code(6400), "06400")
        self.assertEqual(normalize_number("No. 12-A"), "12-A")
        self.assertNotEqual(key_for("Av. Juárez 100"), legacy_key_for("Av. Juárez 100"))
        self.assertEqual(legacy_key_for("Av. Juárez 100"), legacy_key_for("AV. JUAREZ 100"))

    def test_score_uses_all_signals(self):
        google = {"ADDRESS COMPLETE": True, "VALIDATION GRANULARITY": "PREMISE", "GEOCODE GRANULARITY": "PREMISE",
                  "PLACE ID": "x", "LATITUD": 25.0, "LONGITUD": -100.0, "CALLE GOOGLE": "Juárez",
                  "MUNICIPIO GOOGLE": "Monterrey", "CODIGO POSTAL GOOGLE": "64000", "COLONIA GOOGLE": "Centro",
                  "_MISSING": [], "_UNCONFIRMED": []}
        parsed = {"CALLE": "Juarez", "MUNICIPIO": "Monterrey", "CODIGO POSTAL": "64000", "COLONIA": "Centro"}
        sepomex = {"score": 100, "cp": "64000", "municipio": "Monterrey", "colonia": "Monterrey Centro"}
        score, confidence, reasons = score_result(google, parsed, sepomex)
        self.assertEqual((score, confidence), (100, "ALTA")); self.assertGreaterEqual(len(reasons), 7)

    def test_second_pass_is_selective(self):
        good = {"ADDRESS COMPLETE": True, "CALLE GOOGLE": "Juárez", "MUNICIPIO GOOGLE": "Monterrey",
                "CODIGO POSTAL GOOGLE": "64000", "_MISSING": []}
        self.assertFalse(needs_second_pass(good, 90, "ALTA"))
        self.assertTrue(needs_second_pass({**good, "ADDRESS COMPLETE": False}, 70, "MEDIA"))

    def test_consolidation_prefers_official_sepomex(self):
        google = {"CALLE GOOGLE": "Juárez", "NUMERO EXTERIOR GOOGLE": "100", "COLONIA GOOGLE": "Centro",
                  "MUNICIPIO GOOGLE": "Monterrey", "CODIGO POSTAL GOOGLE": "", "ESTADO GOOGLE": "Nuevo León",
                  "ADDRESS COMPLETE": False, "VALIDATION GRANULARITY": "ROUTE", "GEOCODE GRANULARITY": "ROUTE",
                  "_MISSING": ["postal_code"], "_UNCONFIRMED": []}
        parsed = {"CODIGO POSTAL": "", "COLONIA": "Centro", "MUNICIPIO": "Monterrey"}
        sepomex = {"score": 96, "cp": "64000", "colonia": "Monterrey Centro", "municipio": "Monterrey", "estado": "Nuevo León"}
        result = consolidate(google, parsed, sepomex)
        self.assertEqual(result["CODIGO POSTAL OFICIAL"], "64000")
        self.assertIn("SEPOMEX", result["ORIGEN DECISION"])

    def test_sepomex_match(self):
        catalog = [{"d_codigo": "64000", "d_asenta": "Monterrey Centro", "D_mnpio": "Monterrey", "d_estado": "Nuevo León"},
                   {"d_codigo": "67100", "d_asenta": "Guadalupe Centro", "D_mnpio": "Guadalupe", "d_estado": "Nuevo León"}]
        match = best_match(catalog, "Centro de Monterrey", "Monterrey", "64000")
        self.assertEqual(match["cp"], "64000"); self.assertGreater(match["score"], 75)

    def test_cache_round_trip(self):
        with tempfile.TemporaryDirectory() as temp:
            with patch("src.cache.CACHE_DB", Path(temp) / "cache.sqlite3"):
                put("service", "key", "address", 200, {"ok": True})
                self.assertEqual(get("service", "key")["payload"], {"ok": True})

    def test_existing_google_columns_are_never_overwritten(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["DIRECCION"] + GOOGLE_HEADERS)
        sheet.append(["Dirección original"] + [f"ORIGINAL-{index}" for index in range(len(GOOGLE_HEADERS))])
        positions = ensure_headers(sheet)
        write_row(sheet, 2, positions, {header: "REEMPLAZO" for header in GOOGLE_HEADERS})
        for index, header in enumerate(GOOGLE_HEADERS):
            self.assertEqual(sheet.cell(2, positions[header]).value, f"ORIGINAL-{index}")

    def test_existing_table_does_not_receive_overlapping_sheet_filter(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["DIRECCION", "CALLE"])
        sheet.append(["Dirección original", "Calle original"])
        table = Table(displayName="DireccionesFuente", ref="A1:B2")
        sheet.add_table(table)
        positions = ensure_headers(sheet)
        self.assertIsNone(sheet.auto_filter.ref)
        self.assertEqual("C", get_column_letter(positions["CALLE GOOGLE"]))
        synchronized = sheet.tables["DireccionesFuente"]
        self.assertEqual(f"A1:{get_column_letter(sheet.max_column)}2", synchronized.ref)
        self.assertEqual(synchronized.ref, synchronized.autoFilter.ref)
        self.assertEqual(
            list(range(1, len(synchronized.tableColumns) + 1)),
            [column.id for column in synchronized.tableColumns],
        )

    def test_end_to_end_preserves_sources_and_writes_v5(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); input_dir = root / "input"; output_dir = root / "output"; log_dir = root / "logs"
            input_dir.mkdir(); workbook = Workbook(); sheet = workbook.active
            original_headers = ["NOMBRE", "DIRECCION", "CALLE", "NUMERO EXTERIOR", "NUMERO INTERIOR", "COLONIA",
                                "MUNICIPIO", "CODIGO POSTAL", "CONFIANZA PARSEO", "OBSERVACIONES"]
            sheet.append(original_headers); sheet.append(["Persona", "Av. Juárez 100", "Av. Juarez", "100", "", "Centro", "Monterrey", "64000", "ALTA", "Original"])
            source = input_dir / "entrada.xlsx"; workbook.save(source)
            response = {"status": 200, "payload": validation_payload(), "error": None, "cached": False}
            catalog = [{"d_codigo": "64000", "d_asenta": "Monterrey Centro", "D_mnpio": "Monterrey", "d_estado": "Nuevo León"}]
            with patch.multiple(processor, INPUT_DIR=input_dir, OUTPUT_DIR=output_dir, LOG_DIR=log_dir), \
                 patch.object(processor, "validate", return_value=response) as validate_mock, \
                 patch.object(processor, "geocode") as geocode_mock, patch.object(processor, "load_catalog", return_value=catalog):
                output, counts, errors = processor.process()
            self.assertFalse(errors); self.assertEqual(counts["ALTA"], 1); validate_mock.assert_called_once(); geocode_mock.assert_not_called()
            result = load_workbook(output); result_sheet = result.active
            headers = [cell.value for cell in result_sheet[1]]
            for header in OFFICIAL_COLUMNS: self.assertIn(header, headers)
            for index, value in enumerate(["Persona", "Av. Juárez 100", "Av. Juarez", "100", "", "Centro", "Monterrey", "64000", "ALTA", "Original"], 1):
                self.assertEqual(result_sheet.cell(2, index).value or "", value or "")
            decision_log = [json.loads(line) for line in (log_dir / "decisiones_v5.jsonl").read_text(encoding="utf-8").splitlines()]
            self.assertEqual(decision_log[0]["final_confidence"], "ALTA")

if __name__ == "__main__": unittest.main()
