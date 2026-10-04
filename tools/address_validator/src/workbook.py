from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from .consolidation import OFFICIAL_COLUMNS
from .xlsx_integrity import synchronize_tables

GOOGLE_HEADERS = [
    "CALLE GOOGLE", "NUMERO EXTERIOR GOOGLE", "NUMERO INTERIOR GOOGLE", "COLONIA GOOGLE",
    "MUNICIPIO GOOGLE", "CODIGO POSTAL GOOGLE", "ESTADO GOOGLE", "PAIS GOOGLE", "LATITUD",
    "LONGITUD", "PLACE ID", "DIRECCION ESTANDARIZADA GOOGLE", "VALIDATION GRANULARITY",
    "GEOCODE GRANULARITY", "ADDRESS COMPLETE", "CONFIANZA GOOGLE", "OBSERVACIONES GOOGLE",
    "FUENTE FINAL", "CP SEPOMEX", "COLONIA SEPOMEX", "MUNICIPIO SEPOMEX",
    "SIMILITUD SEPOMEX", "VALIDACION POSTAL",
]
NEW_HEADERS = GOOGLE_HEADERS + OFFICIAL_COLUMNS
PROTECTED_PREFIXES = ("DIRECCION", "CALLE GOOGLE", "NUMERO EXTERIOR GOOGLE", "NUMERO INTERIOR GOOGLE",
                      "COLONIA GOOGLE", "MUNICIPIO GOOGLE", "CODIGO POSTAL GOOGLE", "ESTADO GOOGLE",
                      "PAIS GOOGLE", "CONFIANZA GOOGLE", "OBSERVACIONES GOOGLE")

def first_xlsx(input_dir):
    files = sorted(Path(input_dir).glob("*.xlsx"))
    if not files: raise FileNotFoundError("No hay archivo .xlsx dentro de la carpeta input.")
    return files[0]

def open_source(path):
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[0]]
    headers = {str(c.value).strip().upper(): c.column for c in sheet[1] if c.value is not None}
    if "DIRECCION" not in headers: raise ValueError("El Excel no contiene una columna llamada DIRECCION.")
    return workbook, sheet, headers

def ensure_headers(sheet):
    existing = {str(c.value).strip().upper(): c.column for c in sheet[1] if c.value is not None}
    next_column, positions = max(existing.values(), default=0) + 1, {}
    for header in NEW_HEADERS:
        key = header.upper()
        if key in existing:
            positions[header] = existing[key]
        else:
            positions[header] = next_column
            cell = sheet.cell(1, next_column, header)
            cell.fill = PatternFill("solid", fgColor="1F4E78" if header in OFFICIAL_COLUMNS else "385723")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            next_column += 1
    sheet.freeze_panes = "A2"
    synchronize_tables(sheet)
    return positions

def read_context(sheet, row, headers):
    def value(name):
        column = headers.get(name)
        return sheet.cell(row, column).value if column else ""
    return {name: value(name) for name in ("CALLE", "NUMERO EXTERIOR", "NUMERO INTERIOR", "COLONIA",
                                            "MUNICIPIO", "CODIGO POSTAL", "ESTADO")}

def existing_google(sheet, row, headers):
    return {header: sheet.cell(row, headers[header]).value for header in GOOGLE_HEADERS if header in headers}

def write_row(sheet, row, positions, result, preserve_existing=True):
    for header in NEW_HEADERS:
        cell = sheet.cell(row, positions[header])
        if preserve_existing and header in GOOGLE_HEADERS and cell.value not in (None, ""):
            continue
        cell.value = result.get(header, "")

def widths(sheet, positions):
    wide = {"CALLE GOOGLE": 26, "COLONIA GOOGLE": 30, "MUNICIPIO GOOGLE": 24,
            "DIRECCION ESTANDARIZADA GOOGLE": 48, "OBSERVACIONES GOOGLE": 52,
            "COLONIA SEPOMEX": 30, "MUNICIPIO SEPOMEX": 24, "VALIDACION POSTAL": 20,
            "CALLE OFICIAL": 28, "COLONIA OFICIAL": 30, "MUNICIPIO OFICIAL": 24,
            "PLACE ID OFICIAL": 28, "CONFIANZA FINAL": 18,
            "ORIGEN DECISION": 55, "MOTIVO DECISION": 90}
    for header, column in positions.items():
        sheet.column_dimensions[get_column_letter(column)].width = wide.get(header, 17)
    for header in ("CODIGO POSTAL GOOGLE", "CP SEPOMEX", "CODIGO POSTAL OFICIAL"):
        sheet.column_dimensions[get_column_letter(positions[header])].number_format = "@"
