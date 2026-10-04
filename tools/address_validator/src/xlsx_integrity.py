import posixpath
import re
import zipfile
from copy import copy
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import TableColumn

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
TABLE_REL_TYPE = f"{DOC_REL_NS}/table"
CELL_REF_RE = re.compile(r"^\$?([A-Z]{1,3})\$?([1-9][0-9]*):\$?([A-Z]{1,3})\$?([1-9][0-9]*)$")


def _q(namespace, name):
    return f"{{{namespace}}}{name}"


def _last_physical_header_column(sheet):
    columns = [
        cell.column
        for cell in sheet[1]
        if cell.value is not None and str(cell.value).strip()
    ]
    return max(columns, default=0)


def _table_headers(sheet, min_col, max_col, header_row):
    headers = []
    normalized = set()
    for column in range(min_col, max_col + 1):
        value = sheet.cell(header_row, column).value
        name = "" if value is None else str(value)
        if not name.strip():
            raise ValueError(
                f"La tabla no puede incluir el encabezado vacio {get_column_letter(column)}{header_row}."
            )
        key = name.strip().casefold()
        if key in normalized:
            raise ValueError(f"Encabezado de tabla duplicado: {name!r}.")
        normalized.add(key)
        headers.append(name)
    return headers


def synchronize_tables(sheet):
    """Resize and rebuild every existing table from its header row to the last real header."""
    last_header_col = _last_physical_header_column(sheet)
    for table in sheet.tables.values():
        min_col, min_row, _, max_row = range_boundaries(table.ref)
        if min_row != 1:
            raise ValueError(f"La tabla {table.name} no inicia en la fila de encabezados: {table.ref}.")
        if last_header_col < min_col:
            raise ValueError(f"La tabla {table.name} no tiene encabezados fisicos.")
        headers = _table_headers(sheet, min_col, last_header_col, min_row)

        old_by_name = {column.name: column for column in table.tableColumns}
        rebuilt = []
        for identifier, header in enumerate(headers, 1):
            if header in old_by_name:
                column = copy(old_by_name[header])
                column.id = identifier
                column.name = header
            else:
                column = TableColumn(id=identifier, name=header)
            rebuilt.append(column)

        final_ref = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(last_header_col)}{max_row}"
        )
        table.ref = final_ref
        table.tableColumns = rebuilt
        if table.autoFilter is None:
            table.autoFilter = AutoFilter(ref=final_ref)
        else:
            table.autoFilter.ref = final_ref
        # The source sort condition includes the header row while sortState starts at row 2.
        # Rebuilding the table must not serialize that stale, semantically invalid state.
        table.sortState = None


def compact_output_columns(sheet, output_headers):
    """Move an already-calculated output block left across empty formatted columns only."""
    locations = {
        str(cell.value).strip().upper(): cell.column
        for cell in sheet[1]
        if cell.value is not None and str(cell.value).strip()
    }
    output_keys = [header.strip().upper() for header in output_headers]
    output_columns = [locations[key] for key in output_keys if key in locations]
    if not output_columns:
        return
    if len(output_columns) != len(output_headers):
        missing = [header for header, key in zip(output_headers, output_keys) if key not in locations]
        raise ValueError(f"Faltan columnas V5 calculadas: {missing}")
    if output_columns != list(range(min(output_columns), max(output_columns) + 1)):
        raise ValueError("Las columnas V5 calculadas no forman un bloque contiguo.")

    output_set = set(output_keys)
    base_columns = [
        cell.column
        for cell in sheet[1]
        if cell.value is not None
        and str(cell.value).strip()
        and str(cell.value).strip().upper() not in output_set
    ]
    destination = max(base_columns, default=0) + 1
    source = min(output_columns)
    if source == destination:
        return
    if source < destination:
        raise ValueError("El bloque V5 se superpone con los encabezados originales.")
    for column in range(destination, source):
        if sheet.cell(1, column).value not in (None, ""):
            raise ValueError("No se puede compactar: el hueco contiene encabezados reales.")
    source_ref = (
        f"{get_column_letter(source)}1:"
        f"{get_column_letter(max(output_columns))}{sheet.max_row}"
    )
    shift = destination - source
    for source_column in range(source, max(output_columns) + 1):
        destination_column = source_column + shift
        source_letter = get_column_letter(source_column)
        destination_letter = get_column_letter(destination_column)
        sheet.column_dimensions[destination_letter] = copy(sheet.column_dimensions[source_letter])
        sheet.column_dimensions[destination_letter].index = destination_letter
    sheet.move_range(source_ref, cols=destination - source, translate=False)


def _resolve_part(base_part, target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(base_part), target))


def _rels_part(part):
    directory, filename = posixpath.split(part)
    return posixpath.join(directory, "_rels", f"{filename}.rels")


def _relationships(archive, part):
    rels_name = _rels_part(part)
    if rels_name not in archive.namelist():
        return {}, rels_name
    root = ET.fromstring(archive.read(rels_name))
    return {rel.get("Id"): rel for rel in root.findall(_q(PKG_REL_NS, "Relationship"))}, rels_name


def _shared_strings(archive):
    name = "xl/sharedStrings.xml"
    if name not in archive.namelist():
        return []
    root = ET.fromstring(archive.read(name))
    return ["".join(node.text or "" for node in item.iter(_q(MAIN_NS, "t"))) for item in root]


def _cell_value(cell, strings):
    cell_type = cell.get("t")
    if cell_type == "inlineStr":
        inline = cell.find(_q(MAIN_NS, "is"))
        return "" if inline is None else "".join(node.text or "" for node in inline.iter(_q(MAIN_NS, "t")))
    value = cell.find(_q(MAIN_NS, "v"))
    if value is None:
        return ""
    if cell_type == "s":
        return strings[int(value.text)]
    return value.text or ""


def _column_number(cell_ref):
    letters = "".join(character for character in cell_ref if character.isalpha()).upper()
    number = 0
    for character in letters:
        number = number * 26 + ord(character) - 64
    return number


def _width(ref):
    match = CELL_REF_RE.fullmatch(ref or "")
    if not match:
        raise ValueError(f"Rango de tabla invalido: {ref!r}.")
    min_col = _column_number(match.group(1))
    max_col = _column_number(match.group(3))
    if max_col < min_col:
        raise ValueError(f"Rango de tabla invertido: {ref!r}.")
    return min_col, max_col, max_col - min_col + 1


def validate_xlsx_tables(path, require_table=False):
    """Validate table XML plus worksheet/workbook/package relationships."""
    path = Path(path)
    reports = []
    with zipfile.ZipFile(path) as archive:
        damaged_member = archive.testzip()
        if damaged_member:
            raise ValueError(f"Miembro ZIP dañado: {damaged_member}")
        names = set(archive.namelist())
        if require_table and "xl/tables/table1.xml" not in names:
            raise ValueError("Falta xl/tables/table1.xml.")

        # Parse every XML relationship and table-bearing part, not only table1.xml.
        for name in names:
            if name.endswith((".xml", ".rels")):
                ET.fromstring(archive.read(name))

        content_types = ET.fromstring(archive.read("[Content_Types].xml"))
        table_overrides = {
            node.get("PartName", "").lstrip("/")
            for node in content_types
            if node.get("ContentType") == "application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml"
        }

        workbook_part = "xl/workbook.xml"
        workbook = ET.fromstring(archive.read(workbook_part))
        workbook_rels, _ = _relationships(archive, workbook_part)
        strings = _shared_strings(archive)
        for sheet_node in workbook.find(_q(MAIN_NS, "sheets")):
            sheet_rel_id = sheet_node.get(_q(DOC_REL_NS, "id"))
            if sheet_rel_id not in workbook_rels:
                raise ValueError(f"Relacion de hoja inexistente: {sheet_rel_id}")
            sheet_part = _resolve_part(workbook_part, workbook_rels[sheet_rel_id].get("Target"))
            if sheet_part not in names:
                raise ValueError(f"Parte de hoja inexistente: {sheet_part}")
            sheet = ET.fromstring(archive.read(sheet_part))
            sheet_rels, sheet_rels_name = _relationships(archive, sheet_part)
            table_parts = sheet.find(_q(MAIN_NS, "tableParts"))
            if table_parts is None:
                continue
            declared_count = int(table_parts.get("count", "0"))
            actual_parts = table_parts.findall(_q(MAIN_NS, "tablePart"))
            if declared_count != len(actual_parts):
                raise ValueError(f"tableParts count incoherente en {sheet_part}.")

            row1 = sheet.find(f".//{_q(MAIN_NS, 'sheetData')}/{_q(MAIN_NS, 'row')}[@r='1']")
            physical_headers = {}
            if row1 is not None:
                for cell in row1.findall(_q(MAIN_NS, "c")):
                    physical_headers[_column_number(cell.get("r"))] = _cell_value(cell, strings)

            for table_part_node in actual_parts:
                rel_id = table_part_node.get(_q(DOC_REL_NS, "id"))
                rel = sheet_rels.get(rel_id)
                if rel is None or rel.get("Type") != TABLE_REL_TYPE:
                    raise ValueError(f"Relacion de tabla invalida {rel_id} en {sheet_rels_name}.")
                table_part = _resolve_part(sheet_part, rel.get("Target"))
                if table_part not in names:
                    raise ValueError(f"Parte de tabla inexistente: {table_part}")
                if table_part not in table_overrides:
                    raise ValueError(f"Content type de tabla ausente: {table_part}")
                table = ET.fromstring(archive.read(table_part))
                table_ref = table.get("ref")
                min_col, max_col, range_width = _width(table_ref)
                auto_filter = table.find(_q(MAIN_NS, "autoFilter"))
                if auto_filter is None or auto_filter.get("ref") != table_ref:
                    raise ValueError(f"autoFilter inconsistente en {table_part}.")
                columns_node = table.find(_q(MAIN_NS, "tableColumns"))
                if columns_node is None:
                    raise ValueError(f"tableColumns ausente en {table_part}.")
                columns = columns_node.findall(_q(MAIN_NS, "tableColumn"))
                if int(columns_node.get("count", "-1")) != len(columns):
                    raise ValueError(f"tableColumns count incoherente en {table_part}.")
                if len(columns) != range_width:
                    raise ValueError(f"Numero de columnas no coincide con el rango en {table_part}.")
                ids = [int(column.get("id")) for column in columns]
                if ids != list(range(1, len(columns) + 1)):
                    raise ValueError(f"IDs de tableColumn no son consecutivos en {table_part}: {ids}")
                names_in_table = [column.get("name", "") for column in columns]
                if any(not name.strip() for name in names_in_table):
                    raise ValueError(f"Encabezado de tabla vacio en {table_part}.")
                if len({name.strip().casefold() for name in names_in_table}) != len(names_in_table):
                    raise ValueError(f"Encabezado de tabla duplicado en {table_part}.")
                sheet_headers = [physical_headers.get(column, "") for column in range(min_col, max_col + 1)]
                if names_in_table != sheet_headers:
                    raise ValueError(
                        f"Encabezados XML/fisicos discrepantes en {table_part}: "
                        f"{names_in_table!r} != {sheet_headers!r}"
                    )
                reports.append({
                    "sheet": sheet_node.get("name"),
                    "sheet_part": sheet_part,
                    "sheet_rels": sheet_rels_name,
                    "table_part": table_part,
                    "ref": table_ref,
                    "auto_filter_ref": auto_filter.get("ref"),
                    "column_count": len(columns),
                    "headers": names_in_table,
                    "ids": ids,
                })
    if require_table and not reports:
        raise ValueError("El libro no contiene tablas relacionadas con hojas.")
    return reports
